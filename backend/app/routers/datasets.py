import csv
import io
import json
from typing import Any, Dict, Iterable, List

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.ai_service import analyze_dataset
from app.database import get_db
from app.dependencies import current_staff
from app.models import AnalysisResult, DatasetRow, DatasetUpload
from app.schemas import (
    DatasetAnalysisResponse,
    DatasetRowResponse,
    DatasetUploadDetailResponse,
    DatasetUploadResponse,
)

router = APIRouter(prefix="/api/datasets", tags=["Datasets"])

REQUIRED_COLUMNS = {
    "roll_number", "name", "subject", "class", "marks", "max_marks",
    "topic", "assessment", "date",
}


def _read_rows(filename: str, content: bytes) -> Iterable[Dict[str, Any]]:
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "json":
        payload = json.loads(content.decode("utf-8"))
        rows = payload.get("rows") if isinstance(payload, dict) else payload
        if not isinstance(rows, list):
            raise ValueError("JSON must contain an array or an object with a rows array")
        return rows
    if suffix == "csv":
        return csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    if suffix in {"xlsx", "xlsm"}:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(values)]
        except StopIteration as exc:
            raise ValueError("Spreadsheet is empty") from exc
        return [dict(zip(headers, row)) for row in values]
    raise ValueError("Only .xlsx, .xlsm, .csv, and .json files are supported")


def _normalize_row(raw: Dict[str, Any], row_number: int) -> Dict[str, Any]:
    values = {str(key).strip().lower(): value for key, value in raw.items()}
    missing = [column for column in REQUIRED_COLUMNS if column not in values]
    if missing:
        raise ValueError(f"missing columns: {', '.join(sorted(missing))}")
    roll_number = str(values["roll_number"]).strip()
    name = str(values["name"]).strip()
    subject = str(values["subject"]).strip()
    class_name = str(values["class"]).strip()
    if not all((roll_number, name, subject, class_name)):
        raise ValueError("roll_number, name, subject, and class are required")
    marks = float(values["marks"])
    max_marks = float(values["max_marks"])
    if max_marks <= 0 or marks < 0 or marks > max_marks:
        raise ValueError("marks must be between 0 and max_marks")
    return {
        "roll_number": roll_number,
        "student_name": name,
        "subject": subject,
        "class_name": class_name,
        "marks": marks,
        "max_marks": max_marks,
        "topic": str(values["topic"]).strip() if values["topic"] is not None else None,
        "assessment": str(values["assessment"]).strip() if values["assessment"] is not None else None,
        "assessment_date": str(values["date"]).strip() if values["date"] is not None else None,
        "_row_number": row_number,
    }


def _can_access(dataset: DatasetUpload, claims: Dict[str, Any]) -> bool:
    return (
        claims.get("role") == "mentor"
        or (
            claims.get("role") == dataset.uploaded_by_role
            and claims.get("id") == dataset.uploaded_by_id
        )
    )


def _upload_response(dataset: DatasetUpload) -> DatasetUploadResponse:
    return DatasetUploadResponse(
        id=dataset.id,
        filename=dataset.filename,
        file_type=dataset.file_type,
        status=dataset.status,
        row_count=dataset.row_count,
        error_count=dataset.error_count,
        uploaded_at=dataset.uploaded_at,
    )


@router.post("/upload", response_model=DatasetUploadDetailResponse)
async def upload_dataset(
    file: UploadFile = File(...),
    claims: Dict[str, Any] = Depends(current_staff),
    db: Session = Depends(get_db),
):
    if not file.filename:
        raise HTTPException(status_code=400, detail="A dataset file is required")
    try:
        raw_rows = _read_rows(file.filename, await file.read())
        normalized: List[Dict[str, Any]] = []
        errors: List[Dict[str, Any]] = []
        for row_number, raw in enumerate(raw_rows, start=2):
            try:
                normalized.append(_normalize_row(dict(raw), row_number))
            except (TypeError, ValueError) as exc:
                errors.append({"row": row_number, "error": str(exc)})
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError) as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    if not normalized:
        raise HTTPException(
            status_code=422,
            detail={"message": "No valid dataset rows found", "errors": errors},
        )

    suffix = file.filename.lower().rsplit(".", 1)[-1]
    dataset = DatasetUpload(
        filename=file.filename,
        file_type=suffix,
        uploaded_by_id=claims["id"],
        uploaded_by_role=claims["role"],
        status="processed_with_errors" if errors else "processed",
        row_count=len(normalized),
        error_count=len(errors),
    )
    db.add(dataset)
    db.flush()
    for row in normalized:
        row.pop("_row_number", None)
        db.add(DatasetRow(dataset_id=dataset.id, **row))
    db.commit()
    db.refresh(dataset)
    return DatasetUploadDetailResponse(
        **_upload_response(dataset).model_dump(),
        errors=errors,
    )


@router.get("", response_model=List[DatasetUploadResponse])
def list_datasets(
    claims: Dict[str, Any] = Depends(current_staff),
    db: Session = Depends(get_db),
):
    query = db.query(DatasetUpload)
    if claims["role"] != "mentor":
        query = query.filter(
            DatasetUpload.uploaded_by_id == claims["id"],
            DatasetUpload.uploaded_by_role == claims["role"],
        )
    return [_upload_response(dataset) for dataset in query.order_by(DatasetUpload.uploaded_at.desc()).all()]


@router.get("/{dataset_id}", response_model=DatasetUploadResponse)
def get_dataset(
    dataset_id: int,
    claims: Dict[str, Any] = Depends(current_staff),
    db: Session = Depends(get_db),
):
    dataset = db.get(DatasetUpload, dataset_id)
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")
    if not _can_access(dataset, claims):
        raise HTTPException(status_code=403, detail="You cannot access this dataset")
    return _upload_response(dataset)


@router.get("/{dataset_id}/rows", response_model=List[DatasetRowResponse])
def get_dataset_rows(
    dataset_id: int,
    claims: Dict[str, Any] = Depends(current_staff),
    db: Session = Depends(get_db),
):
    dataset = db.get(DatasetUpload, dataset_id)
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")
    if not _can_access(dataset, claims):
        raise HTTPException(status_code=403, detail="You cannot access this dataset")
    return [
        DatasetRowResponse(
            id=row.id, roll_number=row.roll_number, name=row.student_name,
            subject=row.subject, class_name=row.class_name, marks=row.marks,
            max_marks=row.max_marks, topic=row.topic, assessment=row.assessment,
            date=row.assessment_date,
        )
        for row in dataset.rows
    ]


@router.post("/{dataset_id}/analysis", response_model=DatasetAnalysisResponse)
def analyze_uploaded_dataset(
    dataset_id: int,
    claims: Dict[str, Any] = Depends(current_staff),
    db: Session = Depends(get_db),
):
    dataset = db.get(DatasetUpload, dataset_id)
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")
    if not _can_access(dataset, claims):
        raise HTTPException(status_code=403, detail="You cannot analyze this dataset")
    rows = [
        {
            "roll_number": row.roll_number, "name": row.student_name,
            "subject": row.subject, "class": row.class_name, "marks": row.marks,
            "max_marks": row.max_marks, "topic": row.topic,
            "assessment": row.assessment, "date": row.assessment_date,
        }
        for row in dataset.rows
    ]
    try:
        result, model_name = analyze_dataset(rows)
    except RuntimeError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc
    analysis = AnalysisResult(
        dataset_id=dataset.id,
        analysis_type="teacher_dashboard",
        model_name=model_name,
        result_json=json.dumps(result),
    )
    db.add(analysis)
    db.commit()
    db.refresh(analysis)
    return DatasetAnalysisResponse(
        dataset_id=dataset.id, analysis_type=analysis.analysis_type,
        model_name=analysis.model_name, result=result, created_at=analysis.created_at,
    )


@router.get("/{dataset_id}/analysis", response_model=List[DatasetAnalysisResponse])
def list_dataset_analysis(
    dataset_id: int,
    claims: Dict[str, Any] = Depends(current_staff),
    db: Session = Depends(get_db),
):
    dataset = db.get(DatasetUpload, dataset_id)
    if not dataset:
        raise HTTPException(status_code=404, detail="Dataset not found")
    if not _can_access(dataset, claims):
        raise HTTPException(status_code=403, detail="You cannot access this dataset")
    return [
        DatasetAnalysisResponse(
            dataset_id=item.dataset_id, analysis_type=item.analysis_type,
            model_name=item.model_name, result=json.loads(item.result_json),
            created_at=item.created_at,
        )
        for item in dataset.analyses
    ]
