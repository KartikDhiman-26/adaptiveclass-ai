import io
from datetime import datetime, timezone
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from openpyxl import load_workbook
from jose import JWTError
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Class, Mark, Student, Teacher
from app.schemas import MarkImportResponse, MarkInput, MarkResponse
from app.security import decode_access_token

router = APIRouter(prefix="/api/marks", tags=["Marks"])
bearer = HTTPBearer(auto_error=False)


def current_staff(
    credentials: HTTPAuthorizationCredentials = Depends(bearer),
    db: Session = Depends(get_db),
) -> Dict:
    if not credentials:
        raise HTTPException(status_code=401, detail="Bearer token required")
    try:
        claims = decode_access_token(credentials.credentials)
    except JWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    if claims.get("role") not in {"teacher", "mentor"}:
        raise HTTPException(status_code=403, detail="Staff role required")
    return claims


def check_access(cl: Class, subject: str, claims: Dict, db: Session) -> None:
    if claims["role"] == "mentor":
        return
    if cl.teacher_id != claims.get("id"):
        raise HTTPException(status_code=403, detail="Teacher is not assigned to this class")
    if cl.subject.casefold() != subject.casefold():
        raise HTTPException(status_code=403, detail="Teacher is not assigned to this subject")


def validate_mark_input(payload: MarkInput, claims: Dict, db: Session):
    if not payload.subject.strip() or payload.max_marks <= 0:
        raise HTTPException(status_code=422, detail="Subject and positive max_marks are required")
    if payload.marks_obtained < 0 or payload.marks_obtained > payload.max_marks:
        raise HTTPException(status_code=422, detail="marks_obtained must be between 0 and max_marks")
    cl = db.get(Class, payload.class_id)
    student = db.get(Student, payload.student_id)
    if not cl or not student or student.class_id != cl.id:
        raise HTTPException(status_code=404, detail="Class or student not found")
    if cl.subject.casefold() != payload.subject.strip().casefold():
        raise HTTPException(status_code=422, detail="Subject must match the class subject")
    check_access(cl, payload.subject.strip(), claims, db)
    return cl, student


def mark_response(mark: Mark) -> MarkResponse:
    return MarkResponse(
        student_id=mark.student_id, class_id=mark.class_id, subject=mark.subject,
        marks_obtained=mark.marks_obtained, max_marks=mark.max_marks, id=mark.id,
        student_name=mark.student.name, roll_number=mark.student.roll_number,
        percentage=round(mark.marks_obtained / mark.max_marks * 100, 2),
        updated_at=mark.updated_at,
    )


@router.post("", response_model=MarkResponse)
def upsert_mark(payload: MarkInput, claims: Dict = Depends(current_staff),
                db: Session = Depends(get_db)):
    validate_mark_input(payload, claims, db)
    mark = db.query(Mark).filter(
        Mark.student_id == payload.student_id,
        Mark.class_id == payload.class_id,
        Mark.subject == payload.subject.strip(),
    ).first()
    created = mark is None
    if created:
        mark = Mark(student_id=payload.student_id, class_id=payload.class_id,
                    subject=payload.subject.strip())
        db.add(mark)
    mark.marks_obtained = payload.marks_obtained
    mark.max_marks = payload.max_marks
    db.commit()
    db.refresh(mark)
    return mark_response(mark)


@router.get("", response_model=List[MarkResponse])
def list_marks(
    class_id: int,
    subject: Optional[str] = None,
    claims: Dict = Depends(current_staff),
    db: Session = Depends(get_db),
):
    cl = db.get(Class, class_id)
    if not cl:
        raise HTTPException(status_code=404, detail="Class not found")
    requested_subject = (subject or cl.subject).strip()
    check_access(cl, requested_subject, claims, db)
    if requested_subject.casefold() != cl.subject.casefold():
        raise HTTPException(status_code=422, detail="Subject must match the class subject")
    marks = db.query(Mark).filter(
        Mark.class_id == class_id, Mark.subject == requested_subject
    ).order_by(Mark.student_id).all()
    return [mark_response(mark) for mark in marks]


@router.post("/import", response_model=MarkImportResponse)
async def import_marks(
    file: UploadFile = File(...),
    claims: Dict = Depends(current_staff),
    db: Session = Depends(get_db),
):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an .xlsx or .xlsm Excel file")
    try:
        workbook = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value).strip().lower() if value is not None else "" for value in next(rows)]
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read Excel file: {exc}") from exc

    required = {"roll_number", "subject", "marks", "max_marks", "class_id"}
    missing = required - set(headers)
    if missing:
        raise HTTPException(status_code=422, detail=f"Missing columns: {', '.join(sorted(missing))}")
    indexes = {name: headers.index(name) for name in required}
    imported = updated = 0
    errors = []
    for row_number, values in enumerate(rows, start=2):
        try:
            data = {name: values[index] for name, index in indexes.items()}
            cl = db.get(Class, int(data["class_id"]))
            student = db.query(Student).filter(Student.roll_number == str(data["roll_number"]).strip()).first()
            payload = MarkInput(student_id=student.id if student else 0, class_id=int(data["class_id"]),
                                subject=str(data["subject"]).strip(), marks_obtained=float(data["marks"]),
                                max_marks=float(data["max_marks"]))
            validate_mark_input(payload, claims, db)
            mark = db.query(Mark).filter(
                Mark.student_id == payload.student_id, Mark.class_id == payload.class_id,
                Mark.subject == payload.subject,
            ).first()
            if mark:
                updated += 1
            else:
                mark = Mark(student_id=payload.student_id, class_id=payload.class_id,
                            subject=payload.subject)
                db.add(mark)
                imported += 1
            mark.marks_obtained = payload.marks_obtained
            mark.max_marks = payload.max_marks
        except (ValueError, TypeError, HTTPException) as exc:
            errors.append({"row": row_number, "error": exc.detail if isinstance(exc, HTTPException) else str(exc)})
    db.commit()
    return MarkImportResponse(imported=imported, updated=updated, errors=errors)
